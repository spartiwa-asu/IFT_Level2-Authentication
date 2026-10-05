from pathlib import Path
from shutil import copy2
import textwrap
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
out=Path('Assignment instructions/IFT458_Level2_Student_Assignment.docx')
backup=Path('tmp/student-assignment/before-code-walkthrough.docx')
if not backup.exists(): copy2(out,backup)
D=Document(backup)
assets=Path('tmp/student-assignment/code-snapshots');assets.mkdir(exist_ok=True)
font=ImageFont.truetype('/System/Library/Fonts/Menlo.ttc',30)
small=ImageFont.truetype('/System/Library/Fonts/Menlo.ttc',26)
def p(t): D.add_paragraph(t)
def h(t): D.add_heading(t,2)
def page(t):
 q=D.add_heading(t,1);q.paragraph_format.page_break_before=True

def lines(items):
 for label,explanation in items:
  q=D.add_paragraph();q.paragraph_format.space_after=Pt(3)
  q.add_run(label+' ').bold=True;q.add_run(explanation)

def snapshot(file, selected, name):
 source=Path(file).read_text().splitlines(); rows=[]; prev=None
 for n in selected:
  if prev is not None and n!=prev+1 and any(source[i-1].strip() and not source[i-1].lstrip().startswith('//') for i in range(prev+1,n)): rows.append(('...', ''))
  chunks=textwrap.wrap(source[n-1],width=83,replace_whitespace=False,drop_whitespace=False) or ['']
  rows.extend([(str(n) if i==0 else '',t) for i,t in enumerate(chunks)]);prev=n
 im=Image.new('RGB',(1800,100+len(rows)*42+22),'#f7f8fa');draw=ImageDraw.Draw(im)
 draw.text((30,22),file,fill='#17202a',font=small)
 for i,(num,txt) in enumerate(rows):
  draw.text((25,90+i*42),num,fill='#67717f',font=small)
  draw.text((115,90+i*42),txt,fill='#131820',font=font)
 image=assets/(name+'.png');im.save(image)
 q=D.add_paragraph();q.paragraph_format.space_after=Pt(3)
 r=q.add_run();pic=r.add_picture(str(image),width=Inches(6.85));pic._inline.docPr.set('descr','Code extracted from '+file+'. Original line numbers retained. Long lines wrap; ellipses indicate omitted source.')
 q=D.add_paragraph('Source: '+file+'. Original line numbers; long lines wrap. Ellipses mark omitted code; blank lines and comments may be omitted.')
 q.paragraph_format.space_after=Pt(8)
 for r in q.runs:r.italic=True;r.font.size=Pt(9)

page('7 How passwords are hashed before saving')
p('Read this walkthrough alongside the lab. These snapshots are extracted from the supplied source files, not from a database or a live credential. Passwords are hashed with bcrypt, not encrypted: login checks a candidate against the stored hash rather than decrypting a stored password.')
snapshot('models/userModel.js',[1,2,4,*range(49,54),60],'password-hash')
h('Line by line')
lines([
('Line 1','Loads Mongoose, which connects the schema and model to MongoDB.'),
('Line 2','Loads bcryptjs for password hashing and comparison.'),
('Line 4','Converts BCRYPT_ROUNDS from configuration into a number; the expression falls back to 12 when that value is falsy. This is the bcrypt cost setting, not a count of 12 simple hash operations.'),
('Line 49','Registers an asynchronous pre-save hook. Mongoose validation runs before this hook. A normal function gives this access to the user document being saved.'),
('Line 50','Returns early if password was not changed, preventing an unchanged stored hash from being hashed again when other fields are saved.'),
('Line 51','Waits for bcrypt.hash to generate a salted password hash and replaces the document’s password value with that hash. Passing a numeric cost lets bcrypt generate the salt.'),
('Line 52','Removes passwordConfirmation from the document before persistence.'),
('Line 53','Closes the hook function and its registration. Saving can continue after the awaited hash finishes.'),
('Line 60','Compiles and exports the User model. Other modules use this model to create and query user documents.')])
h('What reaches the database')
p('During signup, the request password temporarily exists in application memory. Before the user document is written, this hook replaces it with a hash. The bcrypt output carries information needed for later comparison, including the salt and cost. It is not a reversible encrypted copy of the password.')
p('Library reference: bcryptjs hashing and comparison API, https://github.com/dcodeIO/bcrypt.js')

page('8 How signup saves the user document')
snapshot('controllers/authController.js',[3,6,9,11,12,14,15],'signup-save')
lines([
('Line 3','Imports the User model, including its validation rules and save hook.'),
('Line 6','Defines publicUser, which returns only the user ID, name, and email. The password hash is excluded from this response shape.'),
('Line 9','Exports the asynchronous signup handler called for POST /users/signup.'),
('Line 11','Reads only the four expected fields from the request body. The empty-object fallback avoids destructuring an absent body.'),
('Line 12','Creates and saves a User. Mongoose validates the fields, runs the password save hook, and writes the document to MongoDB. Await resumes after that operation succeeds; a failure prevents the success response.'),
('Line 14','Sends HTTP 201 and the public user profile. Neither a password nor a token is returned by signup.'),
('Line 15','Closes the signup handler.')])
page('Password rules applied before hashing')
snapshot('models/userModel.js',list(range(21,27)),'password-rules')
lines([
('Line 21','Starts the password field definition.'),
('Line 22','Declares a string field.'),
('Line 23','Requires a password and supplies the missing-value message.'),
('Line 24','Requires at least eight characters during schema validation.'),
('Line 25','Excludes password from normal query results. This controls selection; it does not hash the value or guarantee that every response is safe.'),
('Line 26','Closes the password field definition.')])
p('The adjacent passwordConfirmation rule requires confirmation when password changes and compares the two values before hashing. This checks registration input; the next section explains login verification against a stored hash.')
p('Mongoose reference: create triggers save hooks; validation runs before user pre-save hooks. https://mongoosejs.com/docs/middleware.html')

page('9 How login validates the password')
snapshot('controllers/authController.js',[18,19,22,23,24,27,30,31,32],'login-validation')
lines([
('Line 18','Starts the asynchronous login handler.'),
('Line 19','Reads email and password from the JSON request body.'),
('Line 22','Checks whether either required value is missing or falsy.'),
('Line 23','Returns HTTP 400 and stops this handler if either value is absent.'),
('Line 24','Closes the missing-values condition.'),
('Line 27','Converts email into a lowercase string, queries for that user, and explicitly selects the normally hidden password hash. The string conversion prevents an email object from being passed as a MongoDB query operator here.'),
('Line 30','Rejects an unknown user or a failed password comparison. Short-circuit evaluation skips the method call when no user exists. Await waits for the comparison result.'),
('Line 31','Returns the same HTTP 401 message for an unknown email and an incorrect password; this response does not identify which value was wrong.'),
('Line 32','Closes the rejection condition. Only a successful comparison reaches token creation.')])
h('The comparison method called on line 30')
snapshot('models/userModel.js',[56,57,58],'bcrypt-compare')
lines([
('Line 56','Adds isPasswordMatch to User documents and accepts the submitted candidatePassword.'),
('Line 57','Returns bcrypt.compare(candidatePassword, this.password). It uses the stored hash to check the candidate and resolves to true or false. It does not decrypt a password.'),
('Line 58','Closes the document method.')])
p('Schema validation checks whether new input meets rules such as minimum length. Login password verification asks whether the supplied password matches this existing account. These are different checks.')

page('10 How a token is created and returned')
p('This is the continuation of the login handler after the password check succeeds. The file imports randomBytes and createHash from node:crypto and imports the Session model before this code runs.')
snapshot('controllers/authController.js',list(range(34,48)),'token-response')
lines([
('Line 34','Generates 32 cryptographically random bytes and encodes them as a 64-character hexadecimal token. The token is opaque; it is not a JWT or an encoded password.'),
('Line 35','Calculates an expiry 24 hours after the current time. Multiplying by 1000 converts seconds to milliseconds.'),
('Line 36','Begins creating and awaiting a session document in MongoDB.'),
('Line 37','Hashes the random token with SHA-256 and stores the hexadecimal digest as tokenHash. The usable raw token is not stored in this session document.'),
('Line 38','Associates the session with the authenticated user’s MongoDB ID.'),
('Line 39','Adds expiresAt using property shorthand, equivalent to expiresAt: expiresAt.'),
('Line 40','Closes and awaits the Session.create call before sending success.'),
('Line 41','Sets Cache-Control to no-store and prepares an HTTP 200 JSON response. This cache directive is not encryption.'),
('Line 42','Sets the response status label to success.'),
('Line 43','Returns the raw token to the client using property shorthand. The client needs this value for subsequent requests.'),
('Line 44','Returns the expiry timestamp, which JSON serializes as a date string.'),
('Line 45','Returns the public user profile without the password hash.'),
('Line 46','Closes and sends the JSON response.'),
('Line 47','Closes the login handler.')])
p('The password hash and tokenHash have different jobs: bcrypt checks a human-chosen password; SHA-256 identifies a high-entropy random token in the sessions collection. The token is returned in the login response body, not automatically attached to later requests by the server.')

page('11 How the client adds the token to a header')
p('The client must attach the credential to each protected request because HTTP does not automatically remember a previous login. Swagger uses the security configuration below to build that request header after the user selects Authorize.')
snapshot('SwaggerTFT458.js',[69,241],'swagger-bearer')
lines([
('Line 69','Declares bearerAuth as the default security requirement for documented operations. The empty array contains no OAuth scopes; it does not mean anonymous access. Signup and login override this requirement with security: [].'),
('Line 241','Defines bearerAuth as an HTTP bearer scheme. Swagger uses this definition to show the token input and prefix requests with Bearer. The description explains the token lifetime. OpenAPI describes authentication; server middleware actually enforces it.')])
h('The same header in the supplied REST Client example')
snapshot('requests.http',[46,47],'request-header')
lines([
('Line 46','Sends GET to the games endpoint. The host variable is defined earlier in requests.http.'),
('Line 47','In the VS Code REST Client, reads token from the named login request’s JSON response and inserts it after Bearer. Run the named login request first. This templating is a client feature, not JavaScript executed by the server.')])
h('Follow the token from response to request')
for t in ['Execute POST /users/login and obtain token from the successful JSON response.', 'In Swagger, select Authorize, paste only the token, apply it, and close the dialog.', 'Execute GET /games. Swagger adds Authorization: Bearer followed by that token.', 'Read Server response for the real result. Without a valid token, the server returns 401.']:
 D.add_paragraph(t,style='List Bullet')
p('Using the header keeps the credential out of the URL. It does not encrypt it. HTTPS protects headers and bodies in transit; plain HTTP does not. Treat the token as a secret and cover it in screenshots. Use Clear Swagger Data and Reload to clear this tab’s authorization; that action does not invalidate a copied token.')
p('Protocol reference: RFC 6750, sections 2.1 and 5, https://www.rfc-editor.org/rfc/rfc6750')

page('12 How the server verifies the header')
p('The authentication middleware imports createHash, Session, and User. app.js runs this middleware before the game and score routers. This completes the path from password validation to access on the next request.')
snapshot('middleware/authenticate.js',[5,6,7,8,9,10,12,14,15,16,17,18,19,20],'verify-header')
lines([
('Line 5','Exports an async Express middleware function; next passes control onward.'),
('Line 6','Reads Authorization and matches Bearer followed by one space and 64 hexadecimal characters. match[1] contains the token; a missing header becomes an empty string.'),
('Line 7','Defines deny, which sets a Bearer challenge and HTTP 401 before returning JSON.'),
('Line 8','Provides the failure status and login guidance in that JSON body.'),
('Line 9','Closes the deny helper.'),
('Line 10','Immediately rejects a missing or malformed credential.'),
('Line 12','Computes the same SHA-256 token digest used when the session was created.'),
('Line 14','Finds a matching session whose expiry is later than now. Expiry is enforced immediately, independently of delayed database TTL cleanup.'),
('Line 15','Rejects an unknown or expired session.'),
('Line 16','Loads the user referenced by the session.'),
('Line 17','Rejects a token whose user has been deleted.'),
('Line 18','Attaches the verified user document to this request as req.user.'),
('Line 19','Invokes the next middleware or resource handler after all checks pass.'),
('Line 20','Closes the authentication middleware.')])
p('Authentication establishes the caller’s identity. This build does not yet enforce resource ownership or roles; a valid session alone allows the game and score operations described in the lab.')
D.core_properties.subject='Student lab with extracted code snapshots and line by line authentication walkthrough'
D.save(out)
print(out.resolve())
